import html
import re
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt
from fpdf import FPDF

BASE_DIR = Path(__file__).resolve().parent.parent
LOGO_PATH = BASE_DIR / "Image" / "Logo.png"


def sanitize_text(text: str) -> str:
    """Remove typographic/special characters that can break document formatting."""
    replacements = {
        "‘": "'",
        "’": "'",
        "“": '"',
        "”": '"',
        "–": "-",
        "—": "-",
        "…": "...",
        " ": " ",
        "•": "-",
        "®": "(R)",
        "™": "(TM)",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)

    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text).strip()
    return text.encode("latin-1", "replace").decode("latin-1")


def _split_document(text: str):
    lines = [line.strip() for line in sanitize_text(text).splitlines()]
    return [line for line in lines if line]


def _extract_terms(text: str):
    terms = []
    in_terms = False

    for line in _split_document(text):
        upper = line.upper()
        if "TERMS & CONDITIONS" in upper or "TERMS AND CONDITIONS" in upper:
            in_terms = True
            continue

        if in_terms:
            if re.match(r"^(ARTICLE|SECTION|CLAUSE)\s+\d+", upper):
                break
            if line.startswith(("-", "*", "•")):
                value = line.lstrip("-*• ").strip()
                if value:
                    terms.append(value)
            elif re.match(r"^\d+[.)]\s+", line):
                value = re.sub(r"^\d+[.)]\s+", "", line).strip()
                if value:
                    terms.append(value)

    return terms


def format_docx(text: str, doc_type: str) -> bytes:
    """Create a branded Word document with logo, headings, terms table and footer."""
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = document.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(12)

    if LOGO_PATH.exists():
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run()
        run.add_picture(str(LOGO_PATH), width=Inches(1.65))

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(sanitize_text(doc_type).upper())
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)

    for line in _split_document(text):
        if re.match(r"^(ARTICLE|SECTION|CLAUSE)\s+\d+", line, re.I):
            p = document.add_paragraph()
            run = p.add_run(line)
            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
        elif line.startswith(("-", "*", "•")):
            p = document.add_paragraph(style="List Bullet")
            p.add_run(line.lstrip("-*• ").strip())
        else:
            p = document.add_paragraph(line)
            p.paragraph_format.space_after = Pt(6)

    terms = _extract_terms(text)
    if terms:
        heading = document.add_paragraph()
        run = heading.add_run("Terms")
        run.bold = True
        table = document.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"
        table.rows[0].cells[0].text = "No."
        table.rows[0].cells[1].text = "Term"
        for index, term in enumerate(terms, start=1):
            cells = table.add_row().cells
            cells[0].text = str(index)
            cells[1].text = term
            cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("LegalEase - AI-Powered Legal Document Generator").font.name = "Times New Roman"

    output = BytesIO()
    document.save(output)
    return output.getvalue()


class BrandedPDF(FPDF):
    def __init__(self, doc_type: str):
        super().__init__()
        self.doc_type = sanitize_text(doc_type)
        self.set_auto_page_break(auto=True, margin=18)

    def header(self):
        if LOGO_PATH.exists():
            self.image(str(LOGO_PATH), x=85, y=8, w=40)
            self.ln(18)
        self.set_font("Helvetica", "B", 14)
        self.cell(0, 8, self.doc_type.upper(), align="C")
        self.ln(10)

    def footer(self):
        self.set_y(-15)
        self.set_font("Helvetica", "I", 8)
        self.cell(0, 8, "LegalEase - AI-Powered Legal Document Generator", align="C")


def format_pdf(text: str, doc_type: str) -> bytes:
    """Create a branded PDF with logo, bold section headings, bullets and footer."""
    pdf = BrandedPDF(doc_type)
    pdf.add_page()

    for line in _split_document(text):
        if re.match(r"^(ARTICLE|SECTION|CLAUSE)\s+\d+", line, re.I):
            pdf.set_font("Helvetica", "B", 11)
            pdf.multi_cell(0, 7, line)
        elif line.startswith(("-", "*", "•")):
            pdf.set_font("Helvetica", "", 11)
            pdf.multi_cell(0, 6, "- " + line.lstrip("-*• ").strip())
        else:
            pdf.set_font("Helvetica", "", 11)
            pdf.multi_cell(0, 6, line)
        pdf.ln(1)

    out = pdf.output(dest="S")
    if hasattr(out, "encode"):
        return out.encode("latin-1", "replace")
    return bytes(out)


def format_html_preview(text: str) -> str:
    """Convert generated output to styled HTML blocks for Streamlit preview."""
    blocks = []
    for line in _split_document(text):
        safe = html.escape(line)
        if re.match(r"^(ARTICLE|SECTION|CLAUSE)\s+\d+", line, re.I):
            blocks.append(f'<h3>{safe}</h3>')
        elif line.startswith(("-", "*", "•")):
            blocks.append(f'<div class="term">• {html.escape(line.lstrip("-*• ").strip())}</div>')
        else:
            blocks.append(f"<p>{safe}</p>")

    return "\n".join(blocks)
